"""ساخت جزوه‌ی مرتب (PDF یا DOCX) از خلاصه‌های ذخیره‌شده‌ی یه درس."""
import io
import re
from pathlib import Path
from xml.sax.saxutils import escape

import arabic_reshaper
from bidi.algorithm import get_display
try:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError:  # python-docx نصب نیست: فقط خروجی DOCX کار نمی‌کنه، بقیه‌ی برنامه عادی بالا میاد
    Document = None
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

FONTS_DIR = Path(__file__).parent / "fonts"


def _ensure_font(name: str, filename: str) -> None:
    try:
        pdfmetrics.getFont(name)
    except KeyError:
        pdfmetrics.registerFont(TTFont(name, str(FONTS_DIR / filename)))


_ensure_font("Vazir", "Vazirmatn-Regular.ttf")
_ensure_font("Vazir-Bold", "Vazirmatn-Bold.ttf")

DOCX_FONT = "Tahoma"  # روی ویندوز/مک/موبایل هست و حروف فارسی رو پشتیبانی می‌کنه


# ---------- تمیزکاری متن خلاصه (مارک‌داون ساده) ----------
def _blocks(content: str) -> list[tuple[str, str]]:
    """متن رو به بلوک‌های (نوع، متن) تبدیل می‌کنه. نوع: h=تیتر، b=گلوله، p=پاراگراف."""
    blocks: list[tuple[str, str]] = []
    for raw in (content or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        line = line.replace("**", "").replace("__", "").replace("`", "")
        m = re.match(r"^#{1,6}\s*(.+)$", line)
        if m:
            blocks.append(("h", m.group(1).strip()))
            continue
        m = re.match(r"^[-*•]\s+(.+)$", line)
        if m:
            blocks.append(("b", m.group(1).strip()))
            continue
        blocks.append(("p", line))
    return blocks


def _item_title(item: dict) -> str:
    title = (item.get("title") or "").strip()
    return f"{item['number']}. {title}" if title else f"خلاصه‌ی {item['number']}"


# ---------- PDF ----------
def _wrap_fa(text: str, font: str, size: float, max_width: float) -> list[str]:
    """متن فارسی رو (قبل از bidi) بر اساس عرض واقعی به خط‌ها می‌شکنه و هر خط رو جدا برای نمایش آماده می‌کنه."""
    reshaped = arabic_reshaper.reshape(text)
    lines: list[str] = []
    current = ""
    for word in reshaped.split(" "):
        trial = f"{current} {word}" if current else word
        if current and pdfmetrics.stringWidth(trial, font, size) > max_width:
            lines.append(current)
            current = word
        else:
            current = trial
    if current:
        lines.append(current)
    return [get_display(line) for line in lines]


def _footer(canvas, doc) -> None:
    canvas.saveState()
    canvas.setFont("Vazir", 9)
    canvas.setFillColor(colors.HexColor("#888780"))
    canvas.drawCentredString(A4[0] / 2, 10 * mm, str(canvas.getPageNumber()))
    canvas.restoreState()


def build_jozve_pdf_bytes(course_name: str, items: list[dict]) -> bytes:
    if not items:
        raise ValueError("No items to render")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        title=course_name,
    )
    width = doc.width

    course_style = ParagraphStyle(
        "course", fontName="Vazir-Bold", fontSize=20, alignment=TA_RIGHT, leading=28, spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "meta", fontName="Vazir", fontSize=10, alignment=TA_RIGHT,
        textColor=colors.HexColor("#5F5E5A"), spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "item", fontName="Vazir-Bold", fontSize=14, alignment=TA_RIGHT, leading=22,
        textColor=colors.HexColor("#042C53"), spaceBefore=10, spaceAfter=4, keepWithNext=1,
    )
    sub_style = ParagraphStyle(
        "sub", fontName="Vazir-Bold", fontSize=12, alignment=TA_RIGHT, leading=19,
        textColor=colors.HexColor("#173404"), spaceBefore=4, keepWithNext=1,
    )
    body_style = ParagraphStyle(
        "body", fontName="Vazir", fontSize=11, alignment=TA_RIGHT, leading=18,
    )

    def lines_to_paragraphs(text: str, font: str, size: float, style: ParagraphStyle, indent: float = 0):
        out = []
        for line in _wrap_fa(text, font, size, width - indent):
            out.append(Paragraph(escape(line), style))
        return out

    story = []
    story += lines_to_paragraphs(course_name, "Vazir-Bold", 20, course_style)
    story += lines_to_paragraphs(f"{len(items)} خلاصه", "Vazir", 10, meta_style)
    story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#D3D1C7")))

    for item in items:
        story += lines_to_paragraphs(_item_title(item), "Vazir-Bold", 14, heading_style)
        for kind, text in _blocks(item["content"]):
            if kind == "h":
                story += lines_to_paragraphs(text, "Vazir-Bold", 12, sub_style)
            elif kind == "b":
                story += lines_to_paragraphs("• " + text, "Vazir", 11, body_style, indent=6 * mm)
            else:
                story += lines_to_paragraphs(text, "Vazir", 11, body_style)
            story.append(Spacer(1, 3))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=0.4, color=colors.HexColor("#E5E3DB")))

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return buffer.getvalue()


# ---------- DOCX ----------
def _rtl_paragraph(doc, style: str | None = None, space_after: float = 6):
    p = doc.add_paragraph(style=style)
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))  # راست‌به‌چپ؛ تراز پیش‌فرضش خودبه‌خود سمت راسته
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.3
    return p


def _add_rtl_run(p, text: str, size: float, bold: bool = False, color: str | None = None) -> None:
    run = p.add_run(text)
    rPr = run._r.get_or_add_rPr()
    for child in list(rPr):
        rPr.remove(child)
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attr), DOCX_FONT)
    rPr.append(rFonts)
    if bold:
        rPr.append(OxmlElement("w:b"))
        rPr.append(OxmlElement("w:bCs"))
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), str(int(size * 2)))
        rPr.append(el)
    rPr.append(OxmlElement("w:rtl"))


def build_jozve_docx_bytes(course_name: str, items: list[dict]) -> bytes:
    if Document is None:
        raise RuntimeError("python-docx is not installed")
    if not items:
        raise ValueError("No items to render")

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Pt(60)

    p = _rtl_paragraph(doc, style="Title", space_after=4)
    _add_rtl_run(p, course_name, 22, bold=True, color="042C53")
    p = _rtl_paragraph(doc, space_after=14)
    _add_rtl_run(p, f"{len(items)} خلاصه", 10, color="5F5E5A")

    for item in items:
        p = _rtl_paragraph(doc, style="Heading 2", space_after=6)
        _add_rtl_run(p, _item_title(item), 15, bold=True, color="042C53")
        for kind, text in _blocks(item["content"]):
            p = _rtl_paragraph(doc, space_after=4)
            if kind == "h":
                _add_rtl_run(p, text, 12.5, bold=True, color="173404")
            elif kind == "b":
                _add_rtl_run(p, "• " + text, 11.5)
            else:
                _add_rtl_run(p, text, 11.5)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
